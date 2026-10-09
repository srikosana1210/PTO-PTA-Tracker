"""Turns an .xlsx tracker into tests/fixtures/<name>.json (values and formulas per tab) for the Node harness.
Usage: python3 -I tests/build_fixture.py sheet/PTO-PTA-Tracker-2026-27.xlsx tests/fixtures/template.json"""
import json, sys, datetime
import openpyxl

def cell(v):
    if isinstance(v, datetime.datetime): return {'d': v.isoformat()}
    if isinstance(v, datetime.date): return {'d': datetime.datetime(v.year, v.month, v.day).isoformat()}
    return v

src, dst = sys.argv[1], sys.argv[2]
wb = openpyxl.load_workbook(src)
out = []
for ws in wb.worksheets:
    rows = []
    for r in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
        rows.append([cell(c.value) for c in r])
    while rows and all(v is None for v in rows[-1]) and len(rows) > 1 and ws.title not in ('Campus Register',):
        rows.pop()
    out.append({'name': ws.title, 'rows': rows})
json.dump(out, open(dst, 'w'))
print('wrote', dst, [(s['name'], len(s['rows'])) for s in out])
